from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "data" / "output" / "procurement_audit_overview.pptx"

NAVY = RGBColor(11, 18, 32)
INK = RGBColor(30, 41, 59)
SLATE = RGBColor(71, 85, 105)
MUTED = RGBColor(100, 116, 139)
WHITE = RGBColor(248, 250, 252)
BLUE = RGBColor(37, 99, 235)
CYAN = RGBColor(14, 165, 233)
GREEN = RGBColor(22, 163, 74)
AMBER = RGBColor(217, 119, 6)
RED = RGBColor(220, 38, 38)
PALE = RGBColor(241, 245, 249)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def add_bg(slide, color=WHITE):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    slide.shapes._spTree.remove(shape._element)
    slide.shapes._spTree.insert(2, shape._element)


def textbox(slide, text, x, y, w, h, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT, font="Aptos"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def title(slide, heading, subtitle=None, dark=False):
    color = WHITE if dark else INK
    textbox(slide, heading, 0.65, 0.45, 12.0, 0.55, 28, color, True)
    if subtitle:
        textbox(slide, subtitle, 0.68, 1.08, 11.8, 0.4, 12, RGBColor(203, 213, 225) if dark else MUTED)


def footer(slide, number, dark=False):
    color = RGBColor(148, 163, 184) if dark else MUTED
    textbox(slide, "Government Procurement Auditing System | Investigation support prototype", 0.65, 7.12, 9.8, 0.2, 8, color)
    textbox(slide, str(number), 12.3, 7.1, 0.35, 0.2, 9, color, True, PP_ALIGN.RIGHT)


def pill(slide, text, x, y, w, color, text_color=WHITE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.38))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    tf = shape.text_frame
    tf.clear()
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.name = "Aptos"
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = text_color


def card(slide, x, y, w, h, heading, body, accent=BLUE, dark=False):
    fill = RGBColor(17, 24, 39) if dark else WHITE
    line = RGBColor(51, 65, 85) if dark else RGBColor(226, 232, 240)
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h))
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()
    textbox(slide, heading, x + 0.22, y + 0.18, w - 0.4, 0.3, 14, WHITE if dark else INK, True)
    textbox(slide, body, x + 0.22, y + 0.58, w - 0.4, h - 0.72, 11, RGBColor(203, 213, 225) if dark else SLATE)


def metric(slide, x, y, w, value, label, accent=BLUE, dark=False):
    fill = RGBColor(17, 24, 39) if dark else WHITE
    line = RGBColor(51, 65, 85) if dark else RGBColor(226, 232, 240)
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(1.15))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    textbox(slide, value, x + 0.18, y + 0.18, w - 0.35, 0.42, 24, accent, True, font="Aptos Display")
    textbox(slide, label, x + 0.18, y + 0.68, w - 0.35, 0.25, 9, RGBColor(148, 163, 184) if dark else MUTED)


def add_bullets(slide, items, x, y, w, h, size=16, color=INK):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(10)
        p.bullet = True
    return box


# 1. Cover
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, NAVY)
shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.8), 0, Inches(4.6), prs.slide_height)
shape.fill.solid(); shape.fill.fore_color.rgb = BLUE; shape.fill.transparency = 12; shape.line.fill.background()
shape2 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(10.6), 0, Inches(2.7), prs.slide_height)
shape2.fill.solid(); shape2.fill.fore_color.rgb = CYAN; shape2.fill.transparency = 18; shape2.line.fill.background()
pill(slide, "INVESTIGATION SUPPORT", 0.72, 0.8, 2.2, CYAN)
textbox(slide, "Government\nProcurement\nAuditing System", 0.7, 1.55, 7.5, 2.2, 36, WHITE, True, font="Aptos Display")
textbox(slide, "An explainable, evidence-led workflow for identifying procurement cases that deserve human review.", 0.75, 4.25, 6.9, 0.75, 18, RGBColor(203, 213, 225))
textbox(slide, "TED eForms prototype | August 2026 bulk dataset", 0.75, 6.35, 6.9, 0.3, 12, RGBColor(148, 163, 184))
footer(slide, 1, True)

# 2. Why
slide = prs.slides.add_slide(prs.slide_layouts[6]); add_bg(slide); title(slide, "Why this system exists", "Procurement data is large, heterogeneous, and difficult to review consistently."); footer(slide, 2)
card(slide, 0.7, 1.7, 3.8, 2.0, "Scale", "Thousands of notices, lots, bids, organizations, and relationships create a prioritization problem.", BLUE)
card(slide, 4.75, 1.7, 3.8, 2.0, "Signal overload", "A single unusual pattern can be innocent. Investigators need context, not isolated alerts.", AMBER)
card(slide, 8.8, 1.7, 3.8, 2.0, "Trust", "The system must explain its reasoning and avoid declaring any entity corrupt or fraudulent.", GREEN)
textbox(slide, "Core question", 0.75, 4.35, 2.0, 0.3, 13, BLUE, True)
textbox(slide, "Where should a human investigator look first, and what evidence supports that choice?", 0.75, 4.75, 11.6, 0.8, 26, INK, True, font="Aptos Display")

# 3. Guardrails
slide = prs.slides.add_slide(prs.slide_layouts[6]); add_bg(slide, PALE); title(slide, "Design principles and guardrails", "The output is an investigation queue, not an automated verdict."); footer(slide, 3)
card(slide, 0.7, 1.55, 5.85, 1.25, "1. Prioritize, do not accuse", "Scores represent review priority. They are not corruption probabilities.", RED)
card(slide, 6.8, 1.55, 5.85, 1.25, "2. Explain every flag", "Each case carries rule IDs, evidence, thresholds, peer context, and confidence.", BLUE)
card(slide, 0.7, 3.05, 5.85, 1.25, "3. Separate priority from confidence", "A high-priority signal can still have low evidence confidence when context is weak.", AMBER)
card(slide, 6.8, 3.05, 5.85, 1.25, "4. Keep humans in the loop", "Review status and investigator notes are persisted for follow-up and auditability.", GREEN)
textbox(slide, "The system intentionally does not invent ownership, payment trails, or shared-director relationships that are absent from TED.", 0.85, 5.25, 11.5, 0.6, 17, SLATE, True)

# 4. Data scale
slide = prs.slides.add_slide(prs.slide_layouts[6]); add_bg(slide, NAVY); title(slide, "Data processed in the prototype", "Full run against the configured August 2026 TED source folder.", True); footer(slide, 4, True)
metric(slide, 0.75, 1.75, 2.25, "3,753", "successfully parsed notices", CYAN, True)
metric(slide, 3.2, 1.75, 2.25, "26,272", "analytical procurement records", CYAN, True)
metric(slide, 5.65, 1.75, 2.25, "14,380", "procurement lots", CYAN, True)
metric(slide, 8.1, 1.75, 2.25, "19,119", "tenders / bids", CYAN, True)
metric(slide, 10.55, 1.75, 2.0, "21,189", "generated cases", RED, True)
card(slide, 0.75, 3.55, 5.75, 1.7, "Source handling", "The parser reads extracted XML and XML members inside .tar.gz archives, while avoiding duplicate filenames.", CYAN, True)
card(slide, 6.8, 3.55, 5.75, 1.7, "Additional context", "14,871 organizations, 58,067 raw relationships, 104,474 signal records, and 8,809 network nodes were produced.", BLUE, True)
textbox(slide, "Note: 66,281 XML members were seen; 3,753 parsed successfully in the final run.", 0.78, 5.9, 11.5, 0.3, 11, RGBColor(148, 163, 184))

# 5. Architecture
slide = prs.slides.add_slide(prs.slide_layouts[6]); add_bg(slide); title(slide, "End-to-end architecture", "A modular pipeline keeps ingestion, analysis, scoring, and presentation separate."); footer(slide, 5)
steps = [("1", "Ingest", "TED XML /\ntar.gz", CYAN), ("2", "Normalize", "Common tables\n+ data quality", BLUE), ("3", "Context", "Peer groups\n+ entity resolution", GREEN), ("4", "Analyze", "Rules + stats\n+ network + ML", AMBER), ("5", "Prioritize", "Scores + cases\n+ explanations", RED), ("6", "Review", "Dashboard\n+ feedback", BLUE)]
for i, (num, head, body, color) in enumerate(steps):
    x = 0.65 + i * 2.1
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.55), Inches(1.8), Inches(0.65), Inches(0.65))
    circle.fill.solid(); circle.fill.fore_color.rgb = color; circle.line.fill.background()
    textbox(slide, num, x + 0.55, 1.95, 0.65, 0.2, 14, WHITE, True, PP_ALIGN.CENTER)
    textbox(slide, head, x, 2.65, 1.8, 0.3, 13, INK, True, PP_ALIGN.CENTER)
    textbox(slide, body, x, 3.05, 1.8, 0.65, 11, SLATE, False, PP_ALIGN.CENTER)
    if i < len(steps) - 1:
        textbox(slide, "→", x + 1.75, 2.02, 0.4, 0.3, 20, MUTED, True, PP_ALIGN.CENTER)
card(slide, 1.15, 4.65, 11.0, 1.05, "Shared analytical contract", "Every source should map into the same procurement schema so the peer engine, signals, scoring, and dashboard remain source-independent.", BLUE)

# 6. Detection
slide = prs.slides.add_slide(prs.slide_layouts[6]); add_bg(slide, WHITE); title(slide, "How unusual patterns are detected", "Multiple evidence families are fused instead of relying on a single red flag."); footer(slide, 6)
card(slide, 0.7, 1.55, 2.85, 2.0, "Competition", "Single bidder\nLow bidder count\nUnusual bid margin", BLUE)
card(slide, 3.8, 1.55, 2.85, 2.0, "Price / peers", "Deviation from peer median\nIQR outlier\nUnusual contract amount", AMBER)
card(slide, 6.9, 1.55, 2.85, 2.0, "Behavior", "Repeated winner\nBuyer-vendor repetition\nVendor concentration", GREEN)
card(slide, 10.0, 1.55, 2.6, 2.0, "Network / ML", "Repeated co-bidding\nNetwork structure\nAnomaly score", RED)
textbox(slide, "Evidence fusion", 0.75, 4.35, 2.0, 0.3, 13, BLUE, True)
textbox(slide, "rule score + statistical score + behavior score + network score + ML score", 0.75, 4.75, 11.8, 0.55, 23, INK, True, font="Aptos Display")
textbox(slide, "Signals remain inspectable: rule ID, severity, observed value, threshold, and plain-language evidence are retained in each case.", 0.75, 5.55, 11.6, 0.45, 14, SLATE)

# 7. Score bands
slide = prs.slides.add_slide(prs.slide_layouts[6]); add_bg(slide, PALE); title(slide, "Investigation-priority score", "The current bands determine where a human may choose to start review."); footer(slide, 7)
# horizontal band
bands = [("LOW", "0–50", GREEN, 0.0, 5.0), ("MODERATE", ">50–75", AMBER, 5.0, 2.5), ("HIGH", ">75–90", RED, 7.5, 1.5), ("VERY HIGH", ">90–100", RGBColor(127, 29, 29), 9.0, 1.7)]
for label, score, color, x, w in bands:
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8 + x), Inches(2.0), Inches(w), Inches(0.75))
    shape.fill.solid(); shape.fill.fore_color.rgb = color; shape.line.fill.background()
    textbox(slide, label, 0.8 + x, 2.12, w, 0.2, 11, WHITE, True, PP_ALIGN.CENTER)
    textbox(slide, score, 0.8 + x, 2.95, w, 0.3, 12, INK, True, PP_ALIGN.CENTER)
card(slide, 0.9, 4.0, 3.75, 1.35, "Priority is not guilt", "A score says where to look first, not what conclusion to reach.", RED)
card(slide, 4.8, 4.0, 3.75, 1.35, "Confidence is separate", "Peer size, specialized markets, and missing amounts can reduce confidence.", AMBER)
card(slide, 8.7, 4.0, 3.75, 1.35, "Human decision", "Investigators can mark cases new, under review, reviewed, dismissed, or follow-up.", BLUE)

# 8. Dashboard
slide = prs.slides.add_slide(prs.slide_layouts[6]); add_bg(slide, NAVY); title(slide, "Dashboard workflow", "The Streamlit interface turns analytical outputs into an investigator’s work queue.", True); footer(slide, 8, True)
card(slide, 0.75, 1.6, 3.75, 1.6, "Overview", "Dataset scale, flagged cases, score distribution, priority mix, and top procurements.", CYAN, True)
card(slide, 4.8, 1.6, 3.75, 1.6, "Investigation queue", "Filter by priority, score, buyer, and signal; sort cases by review priority.", BLUE, True)
card(slide, 8.85, 1.6, 3.75, 1.6, "Case detail", "See score composition, signal evidence, peer context, explanation, and review controls.", RED, True)
card(slide, 0.75, 3.65, 3.75, 1.6, "Profiles", "Vendor and buyer summaries with participation, concentration, value, and timelines.", GREEN, True)
card(slide, 4.8, 3.65, 3.75, 1.6, "Network", "Explore repeated co-participation and relationship structure with a cautionary interpretation.", AMBER, True)
card(slide, 8.85, 3.65, 3.75, 1.6, "Data quality", "Review parser quality, missingness, and source limitations before acting on a signal.", CYAN, True)
textbox(slide, "Recommended screenshot: Overview page + one Case Detail page showing the score breakdown and evidence cards.", 0.85, 5.95, 11.4, 0.35, 14, RGBColor(203, 213, 225), True)

# 9. Outputs
slide = prs.slides.add_slide(prs.slide_layouts[6]); add_bg(slide); title(slide, "What the system produces", "Outputs are designed for analysis, review, export, and future integration."); footer(slide, 9)
outputs = [
    ("Processed tables", "notices, lots, tenders, organizations, relationships, procurement_data", BLUE),
    ("Case queue", "cases.csv, explanations.csv, investigation_queue.csv", RED),
    ("Analytics", "peer analysis, vendor and buyer summaries, network nodes, signals", GREEN),
    ("Interfaces", "Streamlit dashboard, SQLite database, FastAPI scaffold", CYAN),
]
for i, (head, body, color) in enumerate(outputs):
    y = 1.55 + i * 1.15
    pill(slide, str(i + 1), 0.8, y + 0.05, 0.45, color)
    textbox(slide, head, 1.45, y, 2.4, 0.3, 15, INK, True)
    textbox(slide, body, 3.75, y, 8.3, 0.4, 13, SLATE)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.45), Inches(y + 0.58), Inches(10.6), Inches(0.01))
    line.fill.solid(); line.fill.fore_color.rgb = RGBColor(226, 232, 240); line.line.fill.background()
textbox(slide, "The system is intentionally source-adaptable: an India-specific parser can map into the same common schema later.", 0.85, 6.15, 11.5, 0.35, 14, BLUE, True)

# 10. Limits and roadmap
slide = prs.slides.add_slide(prs.slide_layouts[6]); add_bg(slide, PALE); title(slide, "Current limits and next steps", "This is a strong operational prototype, not a production fraud determination system."); footer(slide, 10)
card(slide, 0.7, 1.55, 5.75, 3.6, "Current limits", "• TED-only prototype source\n• Ownership and payment data are not present\n• Network analysis is evidence of co-participation, not collusion\n• Full-scale runs can be computationally expensive\n• Human review remains essential", AMBER)
card(slide, 6.85, 1.55, 5.75, 3.6, "Next steps", "• Add validated local procurement parsers\n• Add registry and payment joins with provenance\n• Improve sampling and monitoring for scale\n• Add investigator export and audit logs\n• Calibrate thresholds with reviewed outcomes", BLUE)
textbox(slide, "Success criterion", 0.75, 5.65, 2.0, 0.3, 13, GREEN, True)
textbox(slide, "Faster, more consistent human review with transparent evidence and fewer unsupported conclusions.", 0.75, 6.0, 11.7, 0.45, 20, INK, True, font="Aptos Display")

# 11. Screenshot plan
slide = prs.slides.add_slide(prs.slide_layouts[6]); add_bg(slide, NAVY); title(slide, "Screenshots to add", "The deck works without screenshots, but these three visuals will make the demo much more convincing.", True); footer(slide, 11, True)
card(slide, 0.75, 1.55, 3.75, 3.6, "1. Overview", "Capture the dashboard’s metric cards, priority distribution, and top flagged procurements.\n\nPurpose: establish scale and show that the system is operational.", CYAN, True)
card(slide, 4.8, 1.55, 3.75, 3.6, "2. Case Detail", "Capture a case with the score breakdown, risk band, signal evidence, peer context, and disclaimer.\n\nPurpose: prove explainability.", RED, True)
card(slide, 8.85, 1.55, 3.75, 3.6, "3. Network / Queue", "Capture either the filtered investigation queue or the network graph.\n\nPurpose: show workflow depth and the human-review path.", GREEN, True)
textbox(slide, "Best placement: Overview screenshot on the architecture or dashboard slide; Case Detail screenshot on the scoring slide; Queue screenshot on the outputs slide.", 0.85, 5.75, 11.4, 0.55, 15, RGBColor(203, 213, 225), True)

# 12. Close
slide = prs.slides.add_slide(prs.slide_layouts[6]); add_bg(slide, NAVY); footer(slide, 12, True)
pill(slide, "KEY TAKEAWAY", 0.75, 1.0, 1.55, CYAN)
textbox(slide, "Find unusual patterns.\nExplain the evidence.\nLet humans decide.", 0.75, 1.75, 8.7, 2.4, 34, WHITE, True, font="Aptos Display")
textbox(slide, "Government procurement auditing as a transparent prioritization workflow.", 0.8, 4.65, 8.5, 0.45, 18, RGBColor(203, 213, 225))
textbox(slide, "Prototype outputs: 3,753 notices | 26,272 procurements | 21,189 review cases", 0.8, 5.65, 9.0, 0.3, 12, RGBColor(148, 163, 184))

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(OUT)
